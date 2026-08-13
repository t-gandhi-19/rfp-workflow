"""Generate every synthetic fixture, deterministically.

Both the generator and its outputs are committed. Re-running with the same seed
must produce byte-identical JSON and CSV — so everything is sorted, nothing
iterates a set, and no timestamp or random id reaches a file. A test asserts it.

DOCX and PDF are excluded from the byte-identical guarantee: both formats embed
creation metadata. Their *content* parity is guaranteed structurally instead —
one source structure, two renderers — and the extraction parity test proves the
extractor sees the same questions in both.

Usage:  python -m scripts.generate_fixtures [--check]
        --check regenerates into a temporary directory and diffs, without writing.
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import shutil
import sys
import tempfile
from collections import Counter
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

from scripts.fixtures import golden, registry
from scripts.fixtures.corpus import (
    CLOSING_SENTENCES,
    EVIDENCE_SENTENCES,
    GOVERNANCE_SENTENCES,
    PARAPHRASES,
    RISK_SENTENCES,
    TOPICS,
    Topic,
    family_of,
)

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURES = REPO_ROOT / "fixtures"

#: Changing this changes every generated fixture. It is the whole reproducibility
#: contract, so it lives here and nowhere else.
SEED = 20260810

#: Answers must land in this band (build prompt §8).
MIN_WORDS, MAX_WORDS = 150, 300

#: Outcome distribution across the 40 pairs (build prompt §8).
OUTCOME_COUNTS = {"won": 15, "lost": 10, "unknown": 15}

#: Question-type distribution across the same 40 pairs (build prompt §8).
#: Asserted at generation time so a drifted corpus fails with a readable message
#: rather than an IndexError halfway through assigning outcomes.
TYPE_COUNTS = {"technical": 18, "compliance": 8, "company_info": 8, "commercial": 6}

CORPUS_START = date(2023, 1, 15)
CORPUS_END = date(2026, 6, 15)

#: The single confidential pair is bound to this topic and this customer.
CONFIDENTIAL_TOPIC_KEY = "encryption-key-management"


# ---------------------------------------------------------------------------
# Writers that guarantee byte-stability
# ---------------------------------------------------------------------------


def write_json(path: Path, payload: Any) -> None:
    """Sorted keys, fixed separators, trailing newline, LF endings."""
    path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False)
    path.write_text(text + "\n", encoding="utf-8", newline="\n")


def write_csv(path: Path, header: list[str], rows: list[tuple[str, ...]]) -> None:
    """Rows are written in the order given; callers sort before calling."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)


# ---------------------------------------------------------------------------
# Registry
# ---------------------------------------------------------------------------


def generate_registry(root: Path) -> None:
    out = root / "registry"
    write_csv(
        out / "vendors.csv",
        ["code", "name", "services", "active"],
        sorted(tuple(v) for v in registry.VENDORS),
    )
    write_csv(
        out / "products.csv",
        ["code", "name", "description", "domain"],
        sorted(tuple(p) for p in registry.PRODUCTS),
    )
    write_csv(
        out / "certifications.csv",
        ["code", "name", "scope"],
        sorted(tuple(c) for c in registry.CERTIFICATIONS),
    )
    write_csv(
        out / "case_studies.csv",
        ["code", "title", "customer", "domain", "publicly_usable", "summary"],
        sorted(tuple(c) for c in registry.CASE_STUDIES),
    )
    write_csv(
        out / "customers.csv",
        ["id", "name", "industry"],
        sorted(tuple(c) for c in registry.CUSTOMERS),
    )
    write_csv(
        out / "smes.csv",
        ["id", "name", "role", "capabilities"],
        sorted(tuple(s) for s in registry.SMES),
    )
    write_csv(
        out / "capabilities.csv",
        ["id", "name"],
        sorted(tuple(c) for c in registry.CAPABILITIES),
    )
    write_csv(
        out / "locations.csv",
        ["code", "name", "country", "kind"],
        sorted(tuple(loc) for loc in registry.LOCATIONS),
    )


# ---------------------------------------------------------------------------
# Q&A corpus
# ---------------------------------------------------------------------------


def _compose_answer(topic: Topic, index: int, sme: registry.SME) -> tuple[str, list[str]]:
    """Build one answer and the list of registry entities it names."""
    certification = (
        topic.certifications[0]
        if topic.certifications
        else registry.CERTIFICATIONS[index % len(registry.CERTIFICATIONS)].name
    )
    delivery_centres = [loc for loc in registry.LOCATIONS if loc.kind == "delivery centre"]
    location = (
        topic.locations[0]
        if topic.locations
        else delivery_centres[index % len(delivery_centres)].name
    )

    governance = GOVERNANCE_SENTENCES[index % len(GOVERNANCE_SENTENCES)].format(sme_role=sme.role)
    evidence = EVIDENCE_SENTENCES[index % len(EVIDENCE_SENTENCES)].format(
        certification=certification
    )
    closing = CLOSING_SENTENCES[index % len(CLOSING_SENTENCES)].format(location=location)
    risk = RISK_SENTENCES[index % len(RISK_SENTENCES)]

    answer = " ".join([topic.opening, topic.method, risk, governance, evidence, closing])

    cited = sorted(
        {
            *topic.products,
            *topic.vendors,
            *topic.certifications,
            *topic.locations,
            certification,
            location,
        }
    )
    return answer, cited


def _dates(rng: random.Random, count: int) -> list[str]:
    """Dates spread across the corpus window, sorted and unique."""
    span = (CORPUS_END - CORPUS_START).days
    offsets = sorted(rng.sample(range(span), count))
    return [(CORPUS_START + timedelta(days=offset)).isoformat() for offset in offsets]


def generate_qa_pairs() -> list[dict[str, Any]]:
    # A seeded Mersenne Twister is exactly what this needs: reproducibility, not
    # unpredictability. Nothing here is a secret.
    rng = random.Random(SEED)  # noqa: S311
    topics = sorted(TOPICS, key=lambda t: t.key)

    actual_types = dict(sorted(Counter(t.question_type for t in topics).items()))
    expected_types = dict(sorted(TYPE_COUNTS.items()))
    if actual_types != expected_types:
        raise ValueError(f"corpus type distribution is {actual_types}, expected {expected_types}")
    if len(topics) != sum(OUTCOME_COUNTS.values()):
        raise ValueError(
            f"corpus has {len(topics)} topics but the outcome distribution covers "
            f"{sum(OUTCOME_COUNTS.values())}"
        )

    outcomes: list[str] = []
    for value, count in sorted(OUTCOME_COUNTS.items()):
        outcomes.extend([value] * count)
    rng.shuffle(outcomes)

    customers = sorted(registry.CUSTOMERS, key=lambda c: c.id)
    rotation = [c for c in customers if c.name != registry.CONFIDENTIAL_CUSTOMER]
    dates = _dates(rng, len(topics))

    pairs: list[dict[str, Any]] = []
    for index, topic in enumerate(topics):
        capability_id = registry.capability_id_by_name(topic.capability)
        sme = registry.sme_for_capability(capability_id)
        answer_text, cited = _compose_answer(topic, index, sme)

        # The confidential customer is deliberately kept out of the ordinary
        # rotation. If other pairs also belonged to Bluepine, "confidential
        # content never reaches another customer's draft" would stop being
        # cleanly testable — a leak and a legitimate same-customer answer would
        # look identical.
        confidential = topic.key == CONFIDENTIAL_TOPIC_KEY
        customer = (
            registry.CONFIDENTIAL_CUSTOMER if confidential else rotation[index % len(rotation)].name
        )

        pairs.append(
            {
                "id": f"QA-{index + 1:04d}",
                "question_id": f"HQ-{index + 1:04d}",
                "answer_id": f"ANS-{index + 1:04d}",
                "topic_key": topic.key,
                # The SUBJECT, as opposed to the record. `topic_key` is unique and
                # is what the golden source joins through; `topic_family` groups
                # the two halves of a supersession chain, and is what calibration
                # means by "the same question".
                "topic_family": family_of(topic),
                "question": topic.question,
                "answer": answer_text,
                "domain": "cloud_migration",
                "question_type": topic.question_type,
                "capability_id": capability_id,
                "capability": topic.capability,
                "customer": customer,
                "industry": next(c.industry for c in customers if c.name == customer),
                "answer_date": dates[index],
                "author_sme_id": sme.id,
                "outcome": outcomes[index],
                "confidential": confidential,
                "superseded_by": None,
                "cited_entities": cited,
                "word_count": len(answer_text.split()),
            }
        )

    by_key = {pair["topic_key"]: pair for pair in pairs}

    # Wire the supersession chains, and force the superseded answer to be older
    # than the one that replaced it — a chain whose "old" answer is newer would
    # make the recency multiplier and the staleness eval contradict each other.
    for topic in topics:
        if topic.superseded_by is None:
            continue
        old, new = by_key[topic.key], by_key[topic.superseded_by]
        old["superseded_by"] = new["answer_id"]
        if old["answer_date"] >= new["answer_date"]:
            old["answer_date"], new["answer_date"] = new["answer_date"], old["answer_date"]

    for pair in pairs:
        words = pair["word_count"]
        if not MIN_WORDS <= words <= MAX_WORDS:
            raise ValueError(
                f"{pair['id']} ({pair['topic_key']}) is {words} words; "
                f"must be {MIN_WORDS}-{MAX_WORDS}"
            )

    return sorted(pairs, key=lambda p: p["id"])


def generate_paraphrases(pairs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Alternate phrasings of existing corpus questions.

    A paraphrase is a QUESTION, not a Q&A pair. It carries no answer of its own —
    it points at the answer the original question already has, which is what
    makes it safe to add: the corpus still holds 40 answers, so nothing new
    competes for a ranking and the golden expectations are untouched.

    They exist for calibration. The floor is derived from what "genuinely the
    same question" scores, and before these the corpus had nothing to measure
    that on: its only same-subject pairs were four supersession chains with
    byte-identical question text.

    Every family gets the same number of them, so no subject is weighted more
    heavily than another in the resulting statistics.
    """
    families = sorted({pair["topic_family"] for pair in pairs})
    missing = [family for family in families if family not in PARAPHRASES]
    extra = sorted(set(PARAPHRASES) - set(families))
    if missing or extra:
        raise ValueError(
            f"PARAPHRASES must cover exactly the topic families. "
            f"Missing: {missing or 'none'}. Unknown: {extra or 'none'}."
        )

    counts = {len(value) for value in PARAPHRASES.values()}
    if len(counts) != 1:
        raise ValueError(
            f"every family must supply the same number of paraphrases, got sizes {sorted(counts)}; "
            "an uneven count would weight some subjects more heavily in the calibration statistics"
        )

    # The head of a family owns the paraphrases: for a supersession chain that is
    # the v2, because a paraphrase pointing at a retired answer would put text in
    # the index whose only match is something retrieval is supposed to suppress.
    head_by_family: dict[str, dict[str, Any]] = {}
    for pair in sorted(pairs, key=lambda p: p["id"]):
        if pair["superseded_by"] is None:
            head_by_family[pair["topic_family"]] = pair

    rows: list[dict[str, Any]] = []
    for family in families:
        head = head_by_family[family]
        for offset, text in enumerate(PARAPHRASES[family]):
            index = len(rows) + 1
            rows.append(
                {
                    "id": f"QP-{index:04d}",
                    "question_id": f"HQP-{index:04d}",
                    "topic_key": f"{family}-para-{offset + 1}",
                    "topic_family": family,
                    "paraphrase_of": head["question_id"],
                    "answer_id": head["answer_id"],
                    "question": text,
                    "normalized_question": " ".join(text.split()),
                    "domain": head["domain"],
                    "question_type": head["question_type"],
                    "capability_id": head["capability_id"],
                    "customer": head["customer"],
                }
            )

    seen = [row["question"] for row in rows]
    if len(set(seen)) != len(seen):
        raise ValueError("paraphrase question text must be unique across the corpus")
    originals = {pair["question"] for pair in pairs}
    collisions = sorted(originals.intersection(seen))
    if collisions:
        raise ValueError(
            f"paraphrases must not restate an original question verbatim: {collisions}"
        )
    return sorted(rows, key=lambda r: r["id"])


# ---------------------------------------------------------------------------
# Golden RFP
# ---------------------------------------------------------------------------


def golden_source() -> dict[str, Any]:
    """The one structure both renderers consume.

    The injection string is planted here rather than in the question list, so
    the carrier is picked by deterministic rule instead of being hand-assigned.
    The carrier keeps its expected match: a tampered question is still an
    ordinary question, and the system must answer it while refusing to follow
    the instruction buried in it.
    """
    carrier = golden.injection_carrier()
    questions: list[dict[str, Any]] = []
    for question in sorted(golden.QUESTIONS, key=lambda q: q.order):
        is_carrier = question.id == carrier.id
        questions.append(
            {
                "id": question.id,
                "order": question.order,
                "section": question.section,
                "number": question.number,
                "text": (
                    f"{question.text} {golden.INJECTION_STRING}" if is_carrier else question.text
                ),
                "question_type": question.question_type,
                "mandatory": question.mandatory,
                "word_limit": question.word_limit,
                "expects_match": question.expects_match,
                "trap": "injection" if is_carrier else question.trap,
            }
        )

    return {
        "cover": dict(golden.COVER),
        "sections": list(golden.SECTIONS),
        "injection_string": golden.INJECTION_STRING,
        "injection_carrier_question_id": carrier.id,
        "questions": questions,
    }


def annotation(mandatory: bool, word_limit: int | None) -> str:
    """The bracketed suffix both renderers append, and the extractor parses.

    Identical in both renditions so the extractor's field parsing is exercised
    the same way on each.
    """
    parts: list[str] = []
    if mandatory:
        parts.append("Mandatory")
    if word_limit is not None:
        parts.append(f"maximum {word_limit} words")
    return f" [{'; '.join(parts)}]" if parts else ""


def rendered_lines(source: dict[str, Any]) -> list[tuple[str, str]]:
    """(kind, text) pairs describing the document body, renderer-agnostic."""
    lines: list[tuple[str, str]] = []
    for section in source["sections"]:
        lines.append(("section", f"Section {section}"))
        for question in source["questions"]:
            if question["section"] != section:
                continue
            suffix = annotation(question["mandatory"], question["word_limit"])
            lines.append(("question", f"{question['number']} {question['text']}{suffix}"))
    return lines


def render_docx(source: dict[str, Any], path: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    cover = source["cover"]
    title = document.add_paragraph(cover["title"])
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(18)

    for label, key in (
        ("Reference", "reference"),
        ("Issued by", "issuer"),
        ("Issue date", "issued"),
        ("Response deadline", "deadline"),
        ("Contact", "contact"),
    ):
        document.add_paragraph(f"{label}: {cover[key]}")
    document.add_paragraph(cover["preamble"])
    document.add_page_break()

    for kind, text in rendered_lines(source):
        if kind == "section":
            heading = document.add_paragraph(text)
            heading.runs[0].bold = True
            heading.runs[0].font.size = Pt(14)
        else:
            document.add_paragraph(text)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def render_pdf(source: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body", parent=styles["Normal"], fontName="Helvetica", fontSize=11, leading=15
    )
    section_style = ParagraphStyle(
        "SectionHeading", parent=body, fontName="Helvetica-Bold", fontSize=14, spaceBefore=12
    )
    title_style = ParagraphStyle(
        "Title", parent=body, fontName="Helvetica-Bold", fontSize=18, alignment=1, spaceAfter=18
    )

    cover = source["cover"]
    flow: list[Any] = [Paragraph(cover["title"], title_style)]
    for label, key in (
        ("Reference", "reference"),
        ("Issued by", "issuer"),
        ("Issue date", "issued"),
        ("Response deadline", "deadline"),
        ("Contact", "contact"),
    ):
        flow.append(Paragraph(f"{label}: {cover[key]}", body))
    flow.append(Spacer(1, 8))
    flow.append(Paragraph(cover["preamble"], body))
    flow.append(PageBreak())

    for kind, text in rendered_lines(source):
        flow.append(Paragraph(text, section_style if kind == "section" else body))
        flow.append(Spacer(1, 6))

    SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=cover["title"],
        author=cover["issuer"],
    ).build(flow)


def render_response_template(path: Path) -> None:
    """Response template with one slot per question and an SME-TODO block style."""
    document = Document()
    document.styles["Normal"].font.size = Pt(11)

    title = document.add_paragraph("Response to RFP-2026-MIG-014")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(18)
    document.add_paragraph("Prepared for: Meridian Insurance Group")
    document.add_paragraph("Status: DRAFT - not for submission. Human review required.")
    document.add_page_break()

    for section in golden.SECTIONS:
        heading = document.add_paragraph(f"Section {section}")
        heading.runs[0].bold = True
        heading.runs[0].font.size = Pt(14)
        for question in sorted(golden.QUESTIONS, key=lambda q: q.order):
            if question.section != section:
                continue
            label = document.add_paragraph(f"{question.number} {question.text}")
            label.runs[0].italic = True
            # Slot keyed by question order — the assembler fills these by key,
            # never by matching prose.
            document.add_paragraph(f"{{{{answer:{question.order}}}}}")
            document.add_paragraph(f"{{{{sme_todo:{question.order}}}}}")
            document.add_paragraph("")

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


# ---------------------------------------------------------------------------
# Answer key
# ---------------------------------------------------------------------------


def build_answer_key(pairs: list[dict[str, Any]], source: dict[str, Any]) -> dict[str, Any]:
    by_topic = {pair["topic_key"]: pair for pair in pairs}
    superseded_answer_ids = {
        pair["answer_id"] for pair in pairs if pair["superseded_by"] is not None
    }

    questions: list[dict[str, Any]] = []
    for question in source["questions"]:
        topic_key = question["expects_match"]
        expected = by_topic.get(topic_key) if topic_key else None
        heads_chain = bool(
            expected and any(p["superseded_by"] == expected["answer_id"] for p in pairs)
        )
        questions.append(
            {
                "question_id": question["id"],
                "order": question["order"],
                "section": question["section"],
                "question_type": question["question_type"],
                "mandatory": question["mandatory"],
                "word_limit": question["word_limit"],
                "trap": question["trap"],
                "expected_best_match_answer_id": expected["answer_id"] if expected else None,
                "expected_status": "NO_MATCH" if expected is None else "MATCHED",
                "expected_escalation": question["trap"] != "none" or expected is None,
                # A trap names the ONE failure mode that question exists to
                # exercise. Two traps on one question would let a regression in
                # either hide behind the other firing.
                "expected_guardrail": {
                    "unanswerable": "escalate_no_match",
                    "legal": "legal_hard_block",
                    "pricing": "pricing_hard_block",
                    "forbidden_term": "forbidden_term_escalate",
                    "injection": "injection_flagged",
                    "none": None,
                }[question["trap"]],
                "best_match_heads_supersession_chain": heads_chain,
            }
        )

    return {
        "rfp_reference": source["cover"]["reference"],
        "customer": source["cover"]["issuer"],
        "deadline": source["cover"]["deadline"],
        "question_count": len(questions),
        "questions": questions,
        # Named explicitly so the injection eval asserts the flag and pattern on
        # THIS question, rather than inferring "something escalated somewhere".
        "expected_injection_question_id": source["injection_carrier_question_id"],
        "superseded_answer_ids": sorted(superseded_answer_ids),
        "expected_escalation_question_ids": sorted(
            q["question_id"] for q in questions if q["expected_escalation"]
        ),
        "expected_no_match_question_ids": sorted(
            q["question_id"] for q in questions if q["expected_status"] == "NO_MATCH"
        ),
        "supersession_head_question_ids": sorted(
            q["question_id"] for q in questions if q["best_match_heads_supersession_chain"]
        ),
        "counts": {
            "answerable": sum(1 for q in questions if q["trap"] == "none"),
            "unanswerable": sum(1 for q in questions if q["trap"] == "unanswerable"),
            "mandatory": sum(1 for q in questions if q["mandatory"]),
            "with_word_limit": sum(1 for q in questions if q["word_limit"] is not None),
        },
    }


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def generate(root: Path) -> dict[str, Any]:
    generate_registry(root)
    pairs = generate_qa_pairs()
    source = golden_source()

    write_json(root / "qa_pairs.json", pairs)
    write_json(root / "question_paraphrases.json", generate_paraphrases(pairs))
    write_json(root / "golden_rfp_source.json", source)
    write_json(root / "answer_key.json", build_answer_key(pairs, source))

    render_docx(source, root / "golden_rfp.docx")
    render_pdf(source, root / "golden_rfp.pdf")
    render_response_template(root / "templates" / "response_template.docx")

    return {
        "pairs": len(pairs),
        "outcomes": dict(Counter(p["outcome"] for p in pairs)),
        "types": dict(Counter(p["question_type"] for p in pairs)),
        "supersession_chains": sum(1 for p in pairs if p["superseded_by"]),
        "confidential": sum(1 for p in pairs if p["confidential"]),
        "questions": len(source["questions"]),
    }


def check() -> int:
    """Regenerate into a temp directory and diff the byte-stable outputs."""
    stable = [
        "qa_pairs.json",
        "question_paraphrases.json",
        "golden_rfp_source.json",
        "answer_key.json",
        "registry/vendors.csv",
        "registry/products.csv",
        "registry/certifications.csv",
        "registry/case_studies.csv",
        "registry/customers.csv",
        "registry/smes.csv",
        "registry/capabilities.csv",
        "registry/locations.csv",
    ]
    with tempfile.TemporaryDirectory() as tmp:
        scratch = Path(tmp) / "fixtures"
        generate(scratch)
        drifted = [
            name
            for name in stable
            if (FIXTURES / name).read_bytes() != (scratch / name).read_bytes()
        ]
        shutil.rmtree(scratch, ignore_errors=True)
    if drifted:
        sys.stdout.write("fixtures differ from a fresh generation:\n")
        for name in drifted:
            sys.stdout.write(f"  {name}\n")
        return 1
    sys.stdout.write(f"all {len(stable)} byte-stable fixtures match a fresh generation\n")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate synthetic fixtures")
    parser.add_argument(
        "--check",
        action="store_true",
        help="verify committed fixtures match a fresh generation, writing nothing",
    )
    args = parser.parse_args()
    if args.check:
        return check()

    summary = generate(FIXTURES)
    sys.stdout.write(json.dumps(summary, indent=2, sort_keys=True) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
